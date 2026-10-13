"""Import an ICD-10 Excel annex into dbo.Icd10Codes.

Usage: python import_icd10_catalog.py path\to\icd10.xls
Run deploy_icd10_catalog.sql first, or let this script create the table.
"""
from __future__ import annotations

import sys
import unicodedata
from pathlib import Path

import xlrd
from openpyxl import load_workbook
from sqlalchemy import insert, select, text

from app.database import engine
from app.models import Icd10Code


def norm(value: object) -> str:
    value = "" if value is None else str(value).strip().lower()
    value = "".join(c for c in unicodedata.normalize("NFD", value) if unicodedata.category(c) != "Mn")
    return " ".join(value.replace("đ", "d").split())


def read_xls_catalog(path: Path) -> dict[str, str]:
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    try:
        sheet = workbook.sheet_by_name("ICD10")
    except (KeyError, xlrd.XLRDError) as exc:
        raise RuntimeError("Workbook .xls không có sheet ICD10.") from exc
    found: dict[str, str] = {}
    # Official 4469/QĐ-BYT workbook: row 3 onward, column 17 is MÃ BỆNH,
    # column 20 is TÊN BỆNH (zero-based indexes 16 and 19).
    for row_index in range(3, sheet.nrows):
        raw_code, raw_name = sheet.cell_value(row_index, 16), sheet.cell_value(row_index, 19)
        code, name = str(raw_code).strip().upper(), str(raw_name).strip()
        if code and name and code.lower() != "mã bệnh" and len(code) <= 10:
            found[code] = name
    workbook.release_resources()
    return found


def read_xlsx_catalog(path: Path) -> dict[str, str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    found: dict[str, str] = {}
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = None
        code_col = name_col = None
        for row_index, row in enumerate(rows):
            cells = [norm(value) for value in row]
            if header is None:
                if row_index > 40:
                    break
                for index, title in enumerate(cells):
                    if ("ma benh" in title or "ma icd" in title or title in {"ma", "code"}) and code_col is None:
                        code_col = index
                    if ("ten benh" in title or "ten tieng viet" in title or title in {"ten", "name"}) and name_col is None:
                        name_col = index
                if code_col is not None and name_col is not None:
                    header = row_index
                continue
            if code_col >= len(row) or name_col >= len(row):
                continue
            raw_code, raw_name = row[code_col], row[name_col]
            if raw_code is None or raw_name is None:
                continue
            code, name = str(raw_code).strip().upper(), str(raw_name).strip()
            if not code or not name or len(code) > 10:
                continue
            found[code] = name
    workbook.close()
    return found


def read_catalog(path: Path) -> list[dict[str, str]]:
    found = read_xls_catalog(path) if path.suffix.lower() == ".xls" else read_xlsx_catalog(path)
    if len(found) < 10_000:
        raise RuntimeError(f"Only found {len(found)} codes; refusing to replace catalog. Verify the official workbook and header columns.")
    source = "QD4469/2020" if path.suffix.lower() == ".xls" else "TT06/2026/TT-BYT"
    return [{"code": code, "name_vi": name, "source": source} for code, name in sorted(found.items())]


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python import_icd10_catalog.py <official_icd10.xlsx>")
    path = Path(sys.argv[1]).expanduser().resolve()
    if not path.is_file():
        raise SystemExit(f"Workbook not found: {path}")
    records = read_catalog(path)
    Icd10Code.__table__.create(engine, checkfirst=True)
    with engine.connect() as connection:
        existing_codes = set(connection.execute(select(Icd10Code.code)).scalars())
    new_records = [record for record in records if record["code"] not in existing_codes]
    existing_records = [record for record in records if record["code"] in existing_codes]
    insert_rows = [
        {"Code": record["code"], "NameVi": record["name_vi"], "Source": record["source"]}
        for record in new_records
    ]
    with engine.begin() as connection:
        for start in range(0, len(insert_rows), 1000):
            connection.execute(insert(Icd10Code.__table__), insert_rows[start:start + 1000])
        update_sql = text("UPDATE dbo.Icd10Codes SET NameVi = :name_vi, Source = :source WHERE Code = :code")
        for start in range(0, len(existing_records), 1000):
            connection.execute(update_sql, existing_records[start:start + 1000])
    print(f"Imported {len(records):,} ICD-10 codes from {path.name}: {len(new_records):,} inserted, {len(existing_records):,} updated; existing codes not in the workbook were preserved.")


if __name__ == "__main__":
    main()
