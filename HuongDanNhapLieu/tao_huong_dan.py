from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

root = Path(r'C:\PKNK_KSK\HuongDanNhapLieu')
out = root / 'Huong_dan_quy_trinh_nhap_lieu_KSK.docx'

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(91, 105, 112)

def add_step(doc, number, title, purpose, actions, image_name, note=None):
    h = doc.add_heading(f'Bước {number}. {title}', level=1)
    h.runs[0].font.color.rgb = RGBColor(0, 105, 92)
    p = doc.add_paragraph()
    r = p.add_run('Mục đích: '); r.bold = True
    p.add_run(purpose)
    doc.add_paragraph('Thực hiện:', style='List Bullet')
    for action in actions:
        doc.add_paragraph(action, style='List Bullet 2')
    if note:
        p = doc.add_paragraph()
        r = p.add_run('Lưu ý: '); r.bold = True; r.font.color.rgb = RGBColor(180, 83, 9)
        p.add_run(note)
    image = root / image_name
    doc.add_picture(str(image), width=Cm(16.2))
    add_caption(doc, f'Hình {number}. Giao diện {title.lower()}')

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2); section.right_margin = Cm(2)
styles = doc.styles
styles['Normal'].font.name = 'Arial'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); styles['Normal'].font.size = Pt(10.5)
for style_name in ['Title', 'Heading 1', 'Heading 2']:
    styles[style_name].font.name = 'Arial'; styles[style_name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('HƯỚNG DẪN QUY TRÌNH NHẬP LIỆU\nHỆ THỐNG QUẢN LÝ KHÁM SỨC KHỎE'); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(0, 105, 92)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Địa chỉ sử dụng: http://localhost:5500'); r.italic = True; r.font.size = Pt(11)
doc.add_paragraph('')

h = doc.add_heading('Mục đích', level=1); h.runs[0].font.color.rgb = RGBColor(0, 105, 92)
doc.add_paragraph('Tài liệu hướng dẫn cán bộ vận hành tạo đơn vị, tạo gói/đợt khám và nhập hồ sơ khám sức khỏe trên hệ thống KSK.')

h = doc.add_heading('Quy trình tổng quát', level=1); h.runs[0].font.color.rgb = RGBColor(0, 105, 92)
table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, ['Bước', 'Thao tác', 'Kết quả']):
    set_cell_text(cell, text, True, (255,255,255)); shade(cell, '00695C')
for step, action, result in [
    ('1', 'Tạo đơn vị', 'Có công ty/đơn vị và MST để quản lý hồ sơ.'),
    ('2', 'Tạo gói khám / đợt khám', 'Có đợt khám thuộc đúng đơn vị.'),
    ('3', 'Nhập thông tin hồ sơ', 'Tra cứu, lọc và kiểm soát danh sách hồ sơ.'),
    ('4', 'Thêm mới hồ sơ', 'Tạo một bệnh nhân mới cho đơn vị và đợt khám.'),
    ('5', 'Nhập danh sách hồ sơ bằng Excel', 'Nhập nhiều hồ sơ đồng thời cho đợt khám.'),
]:
    cells = table.add_row().cells
    for cell, text in zip(cells, [step, action, result]): set_cell_text(cell, text)

doc.add_page_break()
add_step(doc, 1, 'Tạo đơn vị',
         'Khai báo thông tin công ty/đơn vị tổ chức khám sức khỏe.',
         ['Trên menu, chọn Đơn vị.', 'Chọn Thêm đơn vị.', 'Nhập Tên đơn vị; nhập Mã số thuế (MST) và Mô tả nếu có.', 'Chọn Lưu thông tin.'],
         '01_TaoDonVi.png',
         'MST nên được kiểm tra kỹ vì dùng để liên kết hồ sơ và các đợt khám của đơn vị.')
add_step(doc, 2, 'Tạo gói khám (đợt khám)',
         'Khai báo gói/đợt khám cho một đơn vị đã tạo.',
         ['Trên menu, chọn Đợt khám.', 'Chọn Thêm đợt khám.', 'Chọn Đơn vị, nhập Tên đợt khám, Ngày bắt đầu và Ngày kết thúc.', 'Chọn Lưu.'],
         '02_TaoGoiKham.png',
         'Một đợt khám chỉ thuộc một đơn vị. Cần tạo đợt khám trước khi thêm hoặc import hồ sơ.')
add_step(doc, 3, 'Nhập thông tin hồ sơ',
         'Tra cứu, xem và quản lý danh sách hồ sơ theo đơn vị và đợt khám.',
         ['Trên menu, chọn Danh sách hồ sơ.', 'Chọn Đơn vị, sau đó chọn Đợt khám.', 'Có thể nhập tên, CCCD hoặc điện thoại để thu hẹp kết quả.', 'Chọn Lọc để hiển thị hồ sơ; bấm tên người bệnh để mở và chỉnh sửa.', 'Nếu cần, dùng In Barcode để in MaHS cho máy in nhãn chuyên dụng.'],
         '03_DanhSachHoSo.png',
         'Chức năng Xóa đợt khám chỉ xóa dữ liệu hồ sơ/lần khám của đợt đang chọn; cần xác nhận cẩn thận trước khi thực hiện.')
add_step(doc, 4, 'Thêm mới hồ sơ mới',
         'Tạo một hồ sơ bệnh nhân trực tiếp trên form khám sức khỏe.',
         ['Trên menu, chọn Thêm HS KSK.', 'Kiểm tra hoặc thay đổi Công ty và Đợt khám ở đầu form.', 'Đặt con trỏ vào Quét mã QR CCCD và quét mã; hệ thống tự tách thông tin hành chính.', 'Bổ sung Số điện thoại, địa chỉ, dân tộc, nghề nghiệp và các thông tin còn thiếu.', 'Chọn Lưu thông tin hành chính; tiếp tục nhập các mục 2 đến 6 khi có kết quả khám.'],
         '04_ThemHoSoMoi.png',
         'Hồ sơ mới chỉ được lưu khi đã chọn Công ty, Đợt khám và nhập Họ và tên.')
add_step(doc, 5, 'Thêm mới danh sách hồ sơ bằng Excel',
         'Nhập nhiều hồ sơ hành chính từ tệp Excel vào đúng đơn vị và đợt khám.',
         ['Trên menu, chọn Nhập Excel.', 'Chọn Công ty và Đợt khám cần nhập.', 'Chọn tệp Excel theo mẫu đã chuẩn bị.', 'Kiểm tra dữ liệu xem trước, sau đó chọn Nhập dữ liệu hành chính.', 'Hệ thống thông báo số lượng nhập thành công và chuyển đến Danh sách hồ sơ của đợt khám.'],
         '05_NhapExcel.png',
         'Dữ liệu Họ tên nên để Unicode; cần kiểm tra cột CCCD, ngày sinh, SHS (nếu có) trước khi nhập.')

h = doc.add_heading('Kiểm tra sau nhập', level=1); h.runs[0].font.color.rgb = RGBColor(0, 105, 92)
for text in [
    'Vào Danh sách hồ sơ, chọn đúng Đơn vị và Đợt khám để kiểm tra số lượng.',
    'Mở ngẫu nhiên một số hồ sơ để xác nhận thông tin hành chính, ngày sinh và số CCCD.',
    'Hoàn thiện các mục khám 2–6 trước khi tạo XML hoặc liên thông dữ liệu.',
    'Sao lưu cơ sở dữ liệu định kỳ theo quy định nội bộ.',
]: doc.add_paragraph(text, style='List Bullet')

doc.save(out)
print(out)
